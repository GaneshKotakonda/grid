from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path('/Users/ganeshreddy/Desktop/untitled folder 2')
ASSETS = ROOT / 'report_assets'
OUT = ROOT / 'GridResilience_Project_Report.docx'

def make_trace_figure(path):
    w, h = 1800, 580
    im = Image.new('RGB', (w, h), '#f5f8fc')
    d = ImageDraw.Draw(im)
    try:
        title = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 38)
        body = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 26)
        small = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 21)
    except Exception:
        title = body = small = ImageFont.load_default()
    d.text((65, 35), 'Moderate preset end-to-end test case', fill='#13263c', font=title)
    stages = [
        ('1  CONTINGENCY', 'Line 1 outage', '284.9 MW initial load', '#c99024'),
        ('2  FORECAST', 'Predictive GCN', 'Mean risk 0.8%', '#835bd6'),
        ('3  MITIGATION', 'PPO action', '14.2 MW shed', '#2671ce'),
        ('4  CASCADE', 'Two lines tripped', '0 MW served', '#c64d5b'),
        ('5  RECOVERY', 'No safe actions', 'TOTAL BLACKOUT', '#b4802e'),
    ]
    margin, gap, boxw, top, boxh = 65, 28, 314, 150, 310
    for i, (head, line1, line2, color) in enumerate(stages):
        x = margin + i * (boxw + gap)
        d.rounded_rectangle((x, top, x+boxw, top+boxh), radius=22, fill='white', outline='#cbd6e2', width=3)
        d.rounded_rectangle((x, top, x+boxw, top+18), radius=10, fill=color)
        d.text((x+20, top+45), head, fill='#24364b', font=small)
        d.text((x+20, top+112), line1, fill='#13263c', font=body)
        d.text((x+20, top+178), line2, fill=color, font=small)
        if i < len(stages)-1:
            ax = x+boxw+7
            d.line((ax, top+155, ax+gap-12, top+155), fill='#91a3b7', width=5)
            d.polygon([(ax+gap-12,top+155),(ax+gap-22,top+146),(ax+gap-22,top+164)], fill='#91a3b7')
    d.text((65, 500), 'Run: Line 1 (Bus 1–Bus 5) • load scale 1.10× • capacity 0.020 • seed 42', fill='#516579', font=small)
    im.save(path)

make_trace_figure(ASSETS / 'moderate_test_trace.png')

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
sec.header_distance = Inches(.32)
sec.footer_distance = Inches(.32)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor(36, 50, 66)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 30), ('Heading 1', 19), ('Heading 2', 13), ('Heading 3', 11)]:
    st = styles[name]
    st.font.name = 'Arial'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(8 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = 'GRIDRESILIENCE  |  PROJECT REPORT'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:
    r.font.name = 'Arial'; r.font.size = Pt(8); r.font.color.rgb = RGBColor(92, 108, 126)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = footer.add_run('GridResilience  •  October 2026  •  ')
rr.font.name = 'Arial'; rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(92, 108, 126)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def para(text='', style=None, align=None, boldlead=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold = True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p

def heading(text, level=1):
    return doc.add_heading(text, level=level)

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def table(headers, rows, widths=None, font=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); set_cell_shading(c,'17395E'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            for r in p.runs: r.font.name='Arial'; r.font.size=Pt(font); r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
        set_cell_margins(c)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx%2: set_cell_shading(c,'EDF3F8')
            set_cell_margins(c)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.name='Arial'; r.font.size=Pt(font); r.font.color.rgb=RGBColor(36,50,66)
        if widths:
            for i,width in enumerate(widths): cells[i].width=Inches(width)
    if widths:
        for i,width in enumerate(widths): t.rows[0].cells[i].width=Inches(width)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def caption(text):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(8)
    r=p.add_run(text); r.italic=True; r.font.name='Arial'; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(75,91,108)

def image(path, width=6.75):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(1)
    p.add_run().add_picture(str(path), width=Inches(width))

def page(): doc.add_page_break()

# 1 Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(100)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('GRIDRESILIENCE'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(15); r.font.color.rgb=RGBColor(36,103,174)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(24)
p.add_run('Predictive Control and Recovery\nfor Cascading Grid Failures')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16)
r=p.add_run('Detailed Project Report'); r.font.size=Pt(17); r.font.color.rgb=RGBColor(68,85,105)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70)
p.add_run('Physics-informed forecasting, reinforcement-learning mitigation,\nand safety-constrained service restoration on the IEEE 14-bus test system').font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(78)
p.add_run('Prepared for Academic Evaluation\nOctober 2026').font.size=Pt(11)
page()

# 2 Abstract
heading('Abstract',1)
para('This report presents GridResilience, an end-to-end research platform for studying cascading transmission failures and coordinated response on the IEEE 14-bus test network. The system combines alternating-current (AC) power-flow simulation, deterministic overload propagation, a topology-aware graph convolutional network (GCN) that estimates next-wave line failures, a proximal policy optimization (PPO) controller for preventive actions, and a recovery environment that accepts restoration actions only when power-flow and operating constraints are satisfied. A FastAPI service and React dashboard expose the workflow as inspectable stages rather than a single opaque prediction.')
para('The saved experiments show that the GCN performs strongly on a held-out, scenario-separated synthetic test set: micro F1 is 0.9965 on 1,442 graph samples. Across 300 held-out mitigation scenarios, predictive PPO serves 30.65% of customer load on average, compared with 26.48% for PPO without GCN input, while reducing average shed load from 9.54 MW to 2.91 MW. The post-cascade recovery audit confirms that attempted unsafe actions are rolled back, but it also reveals a substantial performance gap: greedy recovery restores only 0.28 MW on average and leaves the aggregate blackout rate at 74%. A live moderate-preset dashboard run further demonstrates why prediction accuracy and system-level resilience must be assessed separately: the run produced a forecast and a PPO action, yet two lines tripped and the final state was total blackout.')
para('The results establish a meaningful experimental platform and a credible proof of integration. They do not establish field readiness. The evidence is limited to a small benchmark network, synthetic scenarios, saved experimental runs, and a simplified operating model. The report therefore treats the performance results as evidence of research progress and identifies validation, uncertainty reporting, and recovery effectiveness as priorities for the next phase.')
para('Keywords: cascading failure; power systems; graph neural networks; reinforcement learning; preventive control; AC power flow; service restoration.', boldlead='Keywords:')
heading('Report scope',2)
para('The report evaluates the implemented project and its saved artifacts as of October 2026. It covers system design, simulation and learning methods, benchmark results, a single end-to-end test case, interface evidence, limits, and practical next steps. Numerical values are attributed to the repository artifacts listed in the references.')
page()

# 3 Contents and findings
heading('Contents',1)
for line in ['1. Problem definition and objectives','2. System architecture','3. Grid model and cascade simulation','4. Forecasting dataset and GCN evaluation','5. PPO preventive mitigation','6. Safety-constrained recovery','7. End-to-end test case','8. Dashboard and API','9. Implementation and reproducibility','10. Limitations and future work','11. Conclusion','References']:
    para(line)
heading('Principal findings',2)
table(['Evidence area','Result','Interpretation'],[
 ('Forecasting','GCN micro F1 0.9965 on 1,442 held-out graph samples','Strong in-domain predictive fit on generated IEEE-14 scenarios.'),
 ('Mitigation','30.65% average load served with predictive PPO vs 26.48% with PPO without GCN','+4.17 percentage points; 69.5% less shedding.'),
 ('Reliability','Total blackout 68.33% vs 71.33% for PPO without GCN','A 3.00 percentage-point reduction; blackouts remain common.'),
 ('Recovery safety','100% conservation of the reported safety invariant','Unsafe candidate actions are rejected or reverted.'),
 ('Recovery effectiveness','0.28 MW average restored; blackout rate remains 74%','Safety is implemented more successfully than useful restoration.')
], widths=[1.25,2.6,2.9], font=8.5)
para('The central academic contribution is the integration and evaluation of forecasting, control, cascading-failure physics, and constrained restoration in one reproducible workflow. A key qualification is that a high classifier score is not equivalent to high customer-service resilience; the saved and live-run results make that distinction visible.')
page()

# 4 Problem definition
heading('1. Problem definition and objectives',1)
heading('Why cascading failures matter',2)
para('Transmission grids rely on interconnected lines to move power from generators to loads. When a line is removed, AC power flow redistributes across the remaining network. If a remaining line exceeds its thermal limit, it may trip as well, causing further redistribution and additional failures. The process is nonlinear: an initially small disturbance can have a limited effect in one operating state and a much larger effect under high demand or reduced capacity. Operators therefore need tools that connect contingency analysis to timely, physically plausible actions and to safe restoration after an event.')
para('Many research prototypes address only one part of this sequence. A classifier may identify vulnerable elements without modeling corrective action. A controller may act without explicit risk forecasts. A recovery heuristic may recommend reclosure without verifying the resulting AC state. GridResilience is designed to join these stages while keeping the simulated state and each action auditable.')
heading('Project aim',2)
para('The project aims to build and evaluate an integrated research platform that forecasts downstream line failures, selects preventive redispatch or load-shedding actions, simulates overload-driven propagation, and restores service subject to voltage, thermal, generator, and power-flow feasibility constraints.')
heading('Objectives',2)
for t in [
 'Represent the IEEE 14-bus test network and solve AC power flow for changing operating conditions.',
 'Generate reproducible outage and stress scenarios and model secondary trips using an explicit overload threshold.',
 'Train a graph-based model to estimate physical lines that fail in the next cascade wave.',
 'Compare predictive PPO with do-nothing, rule-based, and PPO without GCN observations on held-out scenarios.',
 'Enforce safety checks on candidate recovery actions and audit whether rejected actions can corrupt surviving service.',
 'Expose the complete incident record through a REST API and an interactive dashboard.'
]: bullet(t)
heading('Research questions',2)
para('The evaluation asks three practical questions. First, does explicit grid topology improve next-wave failure forecasting over a non-graph baseline? Second, does giving a PPO mitigation controller GCN predictions improve held-out service outcomes relative to otherwise comparable baselines? Third, can a recovery policy restore service without violating modeled AC constraints or degrading surviving service? The evidence supports positive answers to the first two questions within the generated benchmark, and a safety-oriented but weak effectiveness result for the third.')
page()

# 5 Architecture
heading('2. System architecture',1)
para('GridResilience is organized as a staged cyber-physical workflow. A scenario supplies the initiating outage, load and generation conditions, line-capacity scaling, and random seed. The physical simulator produces bus and line state. The predictive model ranks downstream line risk. PPO selects a sequence of preventive actions. The cascade engine records line trips until the modeled grid stabilizes, blacks out, or fails to converge. Recovery then evaluates candidate reconnection, load-restoration, and redispatch actions on a cloned network and commits only safe states.')
image(ASSETS/'moderate_test_trace.png', 6.8)
caption('Figure 1. Staged trace from the live moderate-preset dashboard run; values are transcribed from the completed incident record.')
heading('Implemented layers',2)
table(['Layer','Main implementation','Output'],[
 ('Network and physics','IEEE-14 loader, AC solver, metric extraction','Bus voltages, line loadings, convergence and served load.'),
 ('Cascade simulation','Overload detection and deterministic wave propagation','Ordered tripping waves and terminal operating status.'),
 ('Forecasting','Graph construction, GCN/MLP/GAT training and inference','Per-line probability of failure in the next wave.'),
 ('Mitigation','Gymnasium state/action environment and PPO controllers','Redispatch or load-shedding action sequence.'),
 ('Recovery','Candidate action preview and safety audit','Committed restoration history and final service.'),
 ('Service and interface','FastAPI endpoints plus React multi-page dashboard','Scenario execution, results, topology, timeline, and benchmarks.')
], widths=[1.2,2.65,2.9], font=8.2)
para('The pipeline keeps the model outputs interpretable at the system level: the API returns inputs, line probabilities, actions, failed lines, restoration outcomes, and an audit log. The dashboard places these outputs on separate pages and preserves the result while navigating within the application. This structure is useful for academic review because it supports inspection of where a run succeeds or fails, instead of reporting only one final score.')
page()

# 6 Grid & cascade
heading('3. Grid model and cascade simulation',1)
heading('IEEE 14-bus representation',2)
para('The physical test case contains 14 buses, 15 transmission lines, and five transformers. Bus features include voltage magnitude and angle, active and reactive load, generation, service status, and whether the bus is supplied. Line loading is measured against a configured limit. The graph representation stores each physical line as two directed edges so that endpoint-specific flow features can be represented while retaining a 15-line physical target space.')
heading('Power-flow and terminal states',2)
para('The simulation uses AC power flow to calculate the operating state after a contingency or control action. A line is marked overloaded when its loading is strictly greater than the configured 100% threshold. If the solver does not converge, the run is classified as NON_CONVERGED. Otherwise, the system can terminate as STABLE, PARTIAL_BLACKOUT, TOTAL_BLACKOUT, or at a configured maximum step count. The blackout classification is based on served load rather than the mere presence of a tripped line.')
heading('Cascade mechanics',2)
para('The cascade engine applies an outage, solves the network, identifies all overloaded in-service lines, and removes that set as a wave. The process repeats deterministically until no overload remains, the maximum number of steps is reached, or the AC solver fails. The ordered waves and the resulting served load are preserved in the trace. This makes the simulated causal sequence explicit but does not represent stochastic relay timing, protection coordination, transient stability, or operator communications.')
heading('Scenario construction',2)
para('The dataset manifest records 5,000 generated scenarios using seed 42. It varies load scale from 0.8 to 1.5, generation scale from 0.9 to 1.1, dispatch and load variation, and line-capacity scale from 0.000005 to 0.02. Up to two initial outages are allowed, with a recorded 20% chance of multiple outages. The unusually broad capacity range creates both moderate and extreme stress regimes; therefore, benchmark performance should be interpreted relative to this scenario generator, not as a direct description of normal utility operating conditions.')
para('The baseline contingency artifact records line loading changes for a selected outage. Such sensitivity checks help demonstrate that the power-flow model redistributes flows across connected elements. They do not by themselves prove that the cascade model reproduces observed historical events. External system validation remains a separate research requirement.')
heading('Evidence artifacts',2)
para('The main implementation is in src/grid_loader.py, src/simulator.py, src/metrics.py, src/cascade.py, and src/scenarios.py. The generated scenario summary and baseline contingency records are retained in outputs/cascade_scenario_summary.csv and outputs/baseline_results.csv.')
page()

# 7 GNN method
heading('4. Forecasting dataset and GCN evaluation',1)
heading('Prediction target and graph features',2)
para('The supervised target is the set of physical lines that fail in the next cascade wave. The current network state is converted to graph tensors with 14 bus nodes and 30 directed edge records. Each bus has eight features; each directed edge has 11 features, including endpoint indices, active and reactive power, line loading, electrical parameters, service status, current limit, and maximum loading threshold. This target horizon is useful for mitigation because the controller can respond before the next simulated wave, while avoiding a vague label such as “eventually fails.”')
heading('Scenario-level split and training controls',2)
para('The generated dataset contains 9,633 graph samples from 5,000 scenarios. The split policy assigns whole scenarios to training, validation, or testing (70%, 15%, and 15%) to reduce leakage between graph states belonging to the same scenario. The stored summary reports 6,725 train, 1,466 validation, and 1,442 test samples. Training uses normalization statistics and positive-class weighting computed from the training data. The classification threshold is selected on validation pooled micro F1 and frozen before test evaluation.')
para('The GCN uses a hidden dimension of 64, dropout 0.2, learning rate 0.001, batch size 64, and early stopping on validation binary cross-entropy with patience 12, up to 80 epochs. The seed is 42. These controls improve reproducibility and make the test score less vulnerable to threshold tuning on the held-out test labels.')
heading('Model comparison results',2)
table(['Model','Micro precision','Micro recall','Micro F1','Exact match','Train time'],[
 ('MLP','0.9473','0.9753','0.9611','0.8939','2.10 s'),
 ('GCN','0.9959','0.9970','0.9965','0.9868','16.49 s'),
 ('GAT','0.9945','0.9974','0.9959','0.9847','48.40 s')
], widths=[.8,1.15,1.1,.95,1.05,1.05], font=8.5)
para('The GCN is the strongest model on the reported test metrics and has the shortest training time of the two graph models. Its micro F1 exceeds the MLP by about 3.54 percentage points. The near-perfect score is compelling evidence that the model learns the generated benchmark’s next-wave labels; it is not evidence of comparable performance on real utility data, under topology changes, or under distribution shift.')
heading('Stress robustness',2)
para('Across the saved stress buckets, the GCN micro F1 ranges from 0.99197 under moderate stress to 0.99835 under extreme stress, while the GAT is close behind. The MLP is weaker, particularly in the moderate bucket. This consistency supports in-generator robustness across the stored capacity regimes. However, the class distribution and physical assumptions are shared with the training generator, and the report does not have confidence intervals or an independently sourced external grid dataset.')
page()

# The single-run forecast values are incorporated into the end-to-end case study.

# 9 RL method
heading('5. PPO preventive mitigation',1)
heading('Environment and observation design',2)
para('The mitigation environment presents an 87-dimensional observation combining bus voltage, active and reactive bus state, line loading, line status, and 15 GCN risk values. This couples physical measurements to predictive information. The action space contains 14 discrete options: no-op; generator increases or decreases in 10 MW increments for four generators; targeted 15% load shedding at selected buses; and uniform 5% shedding. The fixed action catalog makes the control problem tractable and the action trace easy to interpret, although it restricts the controller to a small subset of possible dispatch decisions.')
heading('Reward design',2)
para('The reward combines served load and stability with explicit rewards for preventing waves and reducing overloads. It penalizes line failure, load loss, shedding, excessive action, overload, non-convergence, and blackouts. The stored phase 5.5 configuration assigns a weight of 2 to served load and 10 to stability; terminal penalties are large for non-convergence and total blackout. This objective encodes a tradeoff: shedding can reduce electrical stress but also sacrifices customer service, so a controller must shed selectively and only when expected stability gains justify it.')
heading('Training and held-out evaluation',2)
para('The predictive and no-GCN PPO variants were trained for 30,000 timesteps using four workers and seed 42. The saved configuration records about 2,225 seconds for predictive PPO and 1,327 seconds for PPO without GCN. Both policies were evaluated on the same 300 held-out scenarios alongside do-nothing and rule-based controllers. Matching scenario sets strengthen the comparison; the README and saved CSV present averages, but do not report confidence intervals or significance tests.')
heading('What the GCN contributes',2)
para('The comparison between PPO with and without the GCN is the clearest experiment for estimating the value of predictive input. The predictive policy served more load and shed less. This is consistent with the design hypothesis that anticipating downstream failures can help the agent act more selectively. It does not prove that the GCN alone caused all improvement: a replicated ablation with multiple policy seeds, uncertainty estimates, and identical hyperparameter search budgets would be needed for a stronger causal claim.')
page()

# 10 RL results
heading('Mitigation benchmark results',1)
table(['Controller','Scenarios','Load served','Load shed','Stable rate','Total blackout','Cascade reduction'],[
 ('Do-Nothing','300','27.00%','0.00 MW','14.67%','73.00%','Baseline'),
 ('Rule-Controller','300','26.34%','8.35 MW','16.00%','73.00%','0.86%'),
 ('PPO (No GCN)','300','26.48%','9.54 MW','18.33%','71.33%','2.74%'),
 ('PPO (Predictive GCN)','300','30.65%','2.91 MW','17.67%','68.33%','3.53%')
], widths=[1.45,.65,.8,.8,.8,.9,1.15], font=8.1)
heading('Interpretation',2)
para('The predictive PPO improves average load served by 4.17 percentage points relative to PPO without GCN (30.65% versus 26.48%). Average load shedding falls from 9.54 MW to 2.91 MW, a 69.5% reduction. Total blackout rate falls by 3.00 percentage points relative to PPO without GCN (68.33% versus 71.33%). Against do-nothing, served load increases by 3.65 percentage points, blackout rate falls by 4.67 points, and cascade reduction is 3.53% as defined in the saved comparison table.')
para('The improvement is meaningful but bounded. Total blackout remains the most frequent outcome for every controller, at 68.33% to 73.00%. Predictive PPO does not maximize the stable-grid rate in this table; PPO without GCN has the highest stable rate at 18.33%, while predictive PPO has 17.67%. The results therefore support improved average customer service and reduced blackout incidence against the selected comparators, not broad prevention of cascading events.')
heading('Statistical and reproducibility considerations',2)
para('All four controllers use 300 held-out scenarios, which is a useful matched comparison. The saved table contains point estimates only. Without per-scenario paired differences, confidence intervals, or repeated training seeds, the magnitude and stability of the reported gains cannot be quantified statistically. The project should retain scenario-level outcomes and add bootstrap confidence intervals, paired tests, and sensitivity analysis for reward weights before making stronger general claims.')
image(ASSETS/'benchmarks.png',6.7)
caption('Figure 2. Live benchmark page with the saved 300-scenario mitigation comparison and the phase 6 recovery audit.')
page()

# 11 recovery
heading('6. Safety-constrained recovery',1)
heading('Safety gate',2)
para('The recovery environment begins from the post-cascade network and enumerates candidate actions such as reconnecting a tripped line, restoring load at a bus, or increasing generator output. A candidate is applied to a trial clone, solved with AC power flow, and accepted only if the power flow converges, in-service bus voltages remain between 0.90 and 1.10 per unit, line loading does not exceed 100%, and generator outputs remain within limits. The implementation also checks that committed actions do not reduce served load below the post-cascade baseline. Unsafe candidates are rolled back.')
heading('Audit results',2)
table(['Method','Scenarios','Avg served load','Avg restored','Stable grid','Total blackout','Recovery success'],[
 ('No Recovery','100','29.49 MW / 12.26%','0.00 MW','13%','74%','0%'),
 ('Greedy Recovery','100','29.77 MW / 12.36%','0.28 MW','13%','74%','4%')
], widths=[1.1,.65,1.25,1.0,.8,.9,1.0], font=8.2)
para('The audit supports a safety claim: the reported invariant is conserved across all attempted actions. It does not support a strong effectiveness claim. Greedy recovery adds just 0.28 MW on average, reconnects 0.15 lines, and leaves both the 74% blackout rate and 13% non-convergence rate unchanged. Only 4% of cases are marked recovery successes. The results identify recovery, not forecasting, as the most direct performance bottleneck in the current integrated system.')
heading('Interpretation and design implication',2)
para('A safety layer that refuses unsafe restoration is necessary, but it can become overly conservative when the candidate set, action ordering, initial state, or redispatch options are inadequate. Future work should distinguish “no safe action exists” from “the search policy failed to find a safe action.” Better optimization can increase restoration while retaining the same acceptance checks: for example, rank candidates by served-load gain, explore sequences over a short horizon, and record counterfactual safety outcomes. Any such change should be evaluated against the current invariant audit so that service gains do not silently weaken safety.')

page()
# 12 End-to-end case
heading('7. End-to-end test case',1)
heading('Scenario configuration',2)
table(['Input','Value'],[
 ('Preset','Moderate'),('Initiating outage','Line 1, Bus 1–Bus 5'),('Load scale','1.10×'),('Line-capacity scale','0.020'),('Seed','42'),('Initial load shown by dashboard','284.9 MW')
], widths=[2.25,4.2], font=9)
heading('Observed trace',2)
image(ASSETS/'moderate_test_trace.png',6.8)
caption('Figure 3. Incident stages and measured outcomes transcribed from the live completed dashboard trace.')
para('The run passed through all five visible stages. The GCN risk assessment reported 0.8% mean risk. The PPO emitted one uniform 5% shedding action, displayed as 14.2 MW shed. The post-cascade state showed two tripped lines, zero load served, and 0% service. Recovery committed no actions and returned TOTAL_BLACKOUT. The operations page reported the two failed lines as Line 1 and Line 0. The prediction page ranked Line 5, Line 6, and Line 4 highest at 4.6%, 4.1%, and 1.6%, respectively.')
heading('What this single run demonstrates',2)
para('The case is intentionally informative rather than successful in the everyday sense. It confirms that the integration carries a configured contingency through forecasting, PPO action selection, cascade propagation, and the recovery safety gate, and that the dashboard exposes a complete record. It also demonstrates a material limitation: one preventive action was not enough to avoid a total blackout, and no safe restoration action was committed. A single case cannot estimate overall controller quality; its value is diagnostic. It motivates review of risk calibration, action choice, line trips, and recovery feasibility together.')
para('This run is distinct from the held-out benchmark averages. It should not be used to replace or inflate the 300-scenario results, and it is not an independent validation sample for the trained models.')
page()

# 13 UI screens
heading('8. Dashboard and API',1)
para('The following captured application screens show how the platform presents scenario configuration, grid topology, and saved empirical evidence. The detailed completed-run measurements are documented in the preceding case study and trace figure.')
heading('REST API surface',2)
para('The FastAPI service provides /health for readiness, /grid for the IEEE-14 topology, /simulate to execute a scenario, /results/{run_id} to retrieve a recorded run, and /benchmarks for saved evaluation summaries. These endpoints make the simulator accessible to the React dashboard and allow a separate client to request a run and inspect its structured outputs. The incident record includes prediction, mitigation, cascade, and recovery information so that the workflow can be examined stage by stage.')
image(ASSETS/'operations.png',6.75)
caption('Figure 4. Scenario command deck with outage, load stress, line capacity, and reproducibility seed controls.')
image(ASSETS/'topology.png',6.75)
caption('Figure 5. Interactive IEEE-14 topology view with buses, line identifiers, and playback controls. This capture shows the base topology before the incident was run.')
page()

# Benchmark evidence is presented with the results and the dashboard screenshot section.

# 15 deployment and reproducibility
heading('9. Implementation and reproducibility',1)
heading('Software structure',2)
para('The Python backend is organized into separate modules for grid loading, power-flow simulation, metrics, cascade physics, scenario generation, graph construction, GNN training, reinforcement-learning environments and controllers, recovery checks, API routes, and the end-to-end pipeline. Scripts generate scenarios and datasets, train models, run cascade demonstrations, evaluate recovery, and execute an end-to-end demo. The frontend uses React and Vite and provides route-specific pages over a shared application state.')
heading('Run configuration',2)
table(['Experiment','Stored settings'],[
 ('GCN evaluation','Seed 42; 5,000 generated scenarios; 70/15/15 scenario split; 1,442 test graph samples.'),
 ('PPO comparison','30,000 training timesteps; four workers; seed 42; 300 held-out scenarios.'),
 ('Recovery audit','Seed 42; 100 scenarios; voltage 0.90–1.10 pu; thermal limit 100%; maximum 15 recovery steps.'),
 ('Dashboard case','Moderate preset; Line 1 outage; load 1.10×; capacity scale 0.020; seed 42.')
], widths=[1.55,5.0], font=8.5)
heading('Reproduction and output files',2)
para('The repository README documents dependency installation, frontend build, backend launch, and the development server. Saved artifacts include model checkpoints, training histories, evaluation CSV files, dataset manifests, JSON metric summaries, and recovery action records. To reproduce the reported findings, an evaluator should preserve the current output artifacts, record package versions and platform details, rerun the evaluation scripts from a clean environment, and compare newly generated summaries against the stored CSVs and JSON files.')
para('The project includes automated tests for loaders, graph construction, cascade behavior, dataset generation, training interfaces, RL environments, recovery, API, and UI routes. This report does not claim a fresh full test suite run; the observed verification in this task was the completed live moderate-preset end-to-end scenario and inspection of its interface trace and stored experimental results.')
page()

# 16 limitations and future
heading('10. Limitations and future work',1)
heading('Limitations',2)
for t in [
 'Test-system scale: IEEE-14 is a small benchmark network and cannot capture the operational complexity of a large interconnected utility grid.',
 'Synthetic distribution: model training and testing use scenarios generated from the same repository assumptions; transfer to real dispatch and protection data is unproven.',
 'Simplified dynamics: steady-state AC power flow and overload thresholds omit transient stability, relay coordination, frequency response, communications, and operator behavior.',
 'High residual risk: even predictive PPO reports 68.33% total blackout on its 300-scenario evaluation set.',
 'Weak recovery outcomes: audited greedy recovery leaves the aggregate blackout rate unchanged and restores 0.28 MW on average.',
 'Statistical uncertainty: benchmark tables present point estimates without confidence intervals, per-scenario paired uncertainty, or repeated training seeds.',
 'Forecast calibration: pooled classification metrics do not establish probability calibration, useful warning time, or reliability under distribution shift.',
 'Action-space constraints: fixed redispatch and shedding actions limit the feasible operational response and may contribute to poor recovery.'
]: bullet(t)
heading('Recommended research sequence',2)
table(['Priority','Next step','Evidence to collect'],[
 ('1','Add repeated seeds and paired uncertainty analysis.','Confidence intervals and scenario-level deltas for every controller.'),
 ('2','Improve recovery search and candidate sequencing without changing safety checks.','Load restored, blackout rate, rejected-action reasons, invariant violations.'),
 ('3','Audit GCN calibration and run out-of-distribution tests.','Reliability diagrams, stress-stratified recall, topology and load-shift tests.'),
 ('4','Expand beyond IEEE-14 and document data provenance.','Independent networks and, where available, historical contingency records.'),
 ('5','Add operationally realistic constraints and temporal dynamics.','Generator ramps, reserve, protection timing, transient/frequency measures.')
], widths=[.65,2.75,3.2], font=8.1)
para('The recommended next milestone is not a larger headline accuracy number. It is a more rigorous system-level evaluation: multiple policy seeds, uncertainty bands, an improved recovery optimizer under unchanged safety limits, and a larger test network. This would let the project show whether predictive gains persist when the test bed, scenario distribution, and action opportunities become more demanding.')
page()

# 17 conclusion/references
heading('11. Conclusion',1)
para('GridResilience addresses a consequential power-systems research problem: preventing cascading line failures while preserving customer service and restoring the network safely after damage. Its contribution is a coherent and inspectable pipeline that connects AC physics, next-wave graph forecasting, PPO mitigation, deterministic cascade simulation, constrained recovery, and a working API/dashboard.')
para('The saved benchmark evidence is promising for forecasting and preventive control. The GCN achieves 0.9965 micro F1 on the generated held-out graph set. Predictive PPO outperforms PPO without GCN on average served load, load shedding, and total blackout rate across 300 matched evaluation scenarios. The project also makes a valuable safety distinction in recovery by checking candidate actions against explicit AC operating limits.')
para('At the same time, the current work has not solved cascading outages as an operational problem. The large remaining blackout rates, very limited average recovery, one-run total-blackout case, synthetic-data basis, and small network impose clear limits. The evidence supports describing the project as a strong integrated research prototype with meaningful measured results and a well-defined next research problem: translating predictive accuracy into reliable service preservation and effective, safe recovery.')
heading('References',1)
refs=[
 '[1] GridResilience repository README. System overview, experimental headline results, setup, routes, and API documentation. Project-local file: README.md, accessed 7 October 2026.',
 '[2] GridResilience source modules. src/grid_loader.py, src/simulator.py, src/metrics.py, src/cascade.py, src/scenarios.py, src/graph.py, src/gnn_dataset.py, src/gnn_training.py, src/rl_env.py, src/rl_controllers.py, src/rl_training.py, src/recovery.py, src/pipeline.py, and src/api.py. Project-local implementation, accessed 7 October 2026.',
 '[3] GNN evaluation artifacts. outputs/gnn_test_metrics.json, outputs/gnn_model_comparison.csv, outputs/gnn_stress_robustness.csv, outputs/gnn_dataset_manifest.json, outputs/gnn_dataset_summary.csv, and outputs/gnn_decision_thresholds.json. Project-local saved results.',
 '[4] Reinforcement-learning evaluation artifacts. outputs/rl_method_comparison.csv, outputs/rl_config.json, and outputs/rl_evaluation.csv. Project-local saved results.',
 '[5] Recovery evaluation artifacts. outputs/recovery_summary.csv, outputs/recovery_evaluation.csv, outputs/recovery_actions.csv, and outputs/recovery_config.json. Project-local saved results.',
 '[6] Cascade evaluation artifacts. outputs/cascade_scenario_summary.csv, outputs/cascade_steps.csv, and outputs/baseline_results.csv. Project-local saved results.',
 '[7] GridResilience dashboard screenshots. Captured from the local React application at the operations, topology, and benchmarks routes on 7 October 2026. Figures 2, 4, and 5 show saved empirical evidence, scenario configuration, and the base topology; the completed moderate test case is transcribed separately in Figure 3 and Section 7.'
]
for ref in refs:
    p=doc.add_paragraph(ref); p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.first_line_indent=Inches(-.2); p.paragraph_format.space_after=Pt(6)

doc.core_properties.title='GridResilience Predictive Control and Recovery Project Report'
doc.core_properties.subject='Academic evaluation of cascading grid failure forecasting, mitigation, and recovery'
doc.core_properties.keywords='GridResilience, power systems, GCN, PPO, cascading failure, recovery'
doc.save(OUT)
print(OUT)
